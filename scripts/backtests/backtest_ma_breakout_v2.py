"""
单均线突破策略 v2（优化版）
- 入场：基础突破 + 5 项过滤（长周期趋势 / 斜率 / ATR 幅度 / 成交量 / 回撤后突破）
- 止损：入场 K 线最低点（保底硬止损）
- 移动出场：ATR 移动止损 + Chandelier Exit（取更紧者）
- 成本：手续费 0.10% + 滑点 0.05%（单边）
- 品种：BTC、ETH、美股市值前 10、港股市值前 5
- 输出：Excel（各参数表 / 最优排名 / 稳定性热力图 / 各品种逐笔明细）
"""

import datetime
import os
import yfinance as yf
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# ── 品种 ──────────────────────────────────────────────────────────────────────
SYMBOLS = {
    # 加密
    "BTC-USD": "Bitcoin",
    "ETH-USD": "Ethereum",
    # 美股市值前 10
    "NVDA":    "NVIDIA",
    "MSFT":    "Microsoft",
    "AAPL":    "Apple",
    "GOOGL":   "Alphabet",
    "AMZN":    "Amazon",
    "META":    "Meta",
    "AVGO":    "Broadcom",
    "TSLA":    "Tesla",
    "BRK-B":   "Berkshire",
    "JPM":     "JPMorgan",
    # 港股市值前 5
    "0700.HK": "腾讯控股",
    "9988.HK": "阿里巴巴",
    "0941.HK": "中国移动",
    "1299.HK": "友邦保险",
    "0939.HK": "建设银行",
}

# ── 交易成本 ──────────────────────────────────────────────────────────────────
FEE_RATE = 0.001
SLIPPAGE = 0.0005
COST     = FEE_RATE + SLIPPAGE     # 单边 0.15%

# ── 参数 ──────────────────────────────────────────────────────────────────────
MA_PERIODS    = list(range(10, 210, 10))
ATR_LEN       = 14
TREND_LEN     = 200
TREND_FALLBACK = 120
SLOPE_LOOKBACK = 5
ATR_BREAKOUT_K = 0.3
VOL_MA_LEN     = 20
VOL_MULT       = 1.2
PULLBACK_LEN   = 5
TRAIL_ATR_K    = 3.0     # ATR 移动止损倍数
CHANDELIER_K   = 3.0     # Chandelier Exit 倍数

END_DATE = datetime.date.today().isoformat()


# ── 指标 ──────────────────────────────────────────────────────────────────────

def compute_sma(values, n):
    out = [None] * len(values)
    s = 0.0
    for i, v in enumerate(values):
        s += v
        if i >= n:
            s -= values[i - n]
        if i >= n - 1:
            out[i] = s / n
    return out


def compute_atr(highs, lows, closes, n=14):
    """Wilder's ATR"""
    out = [None] * len(closes)
    if len(closes) < n + 1:
        return out
    trs = [None] * len(closes)
    trs[0] = highs[0] - lows[0]
    for i in range(1, len(closes)):
        tr = max(
            highs[i] - lows[i],
            abs(highs[i] - closes[i - 1]),
            abs(lows[i]  - closes[i - 1]),
        )
        trs[i] = tr
    # Wilder seed = simple average of first n TRs
    seed = sum(trs[1:n + 1]) / n
    out[n] = seed
    for i in range(n + 1, len(closes)):
        out[i] = (out[i - 1] * (n - 1) + trs[i]) / n
    return out


# ── 入场过滤 ──────────────────────────────────────────────────────────────────

def passes_entry_filters(i, closes, ma, ma_long, atr, vol, vol_ma, has_volume):
    """6 项过滤（基础信号 + 5 项过滤）。i 必须是已经满足基础突破的下标。"""
    # 2. 长周期趋势
    if ma_long[i] is None or closes[i] <= ma_long[i]:
        return False
    # 3. MA 斜率
    j = i - SLOPE_LOOKBACK
    if j < 0 or ma[j] is None or ma[i] <= ma[j]:
        return False
    # 4. ATR 突破幅度
    if atr[i] is None or (closes[i] - ma[i]) < ATR_BREAKOUT_K * atr[i]:
        return False
    # 5. 成交量
    if has_volume:
        if vol_ma[i] is None or vol[i] < VOL_MULT * vol_ma[i]:
            return False
    # 6. 回撤后突破：前 PULLBACK_LEN 根 close 至少有一次 < MA
    pulled_back = False
    for k in range(max(0, i - PULLBACK_LEN), i):
        if ma[k] is not None and closes[k] < ma[k]:
            pulled_back = True
            break
    if not pulled_back:
        return False
    return True


# ── 单品种单参数回测 ──────────────────────────────────────────────────────────

def run_single(highs, lows, closes, volumes, dates, n, has_volume):
    """
    返回字典；trade_log 元素：
      (entry_idx, exit_idx, ret, exit_reason, hard_stop, exit_px_raw, rr)
    """
    if len(closes) <= max(n, TREND_FALLBACK) + ATR_LEN + 2:
        return None

    ma     = compute_sma(closes, n)
    ma_t1  = compute_sma(closes, TREND_LEN)
    ma_t2  = compute_sma(closes, TREND_FALLBACK)
    ma_long = [ma_t1[i] if ma_t1[i] is not None else ma_t2[i] for i in range(len(closes))]
    atr    = compute_atr(highs, lows, closes, ATR_LEN)
    vol_ma = compute_sma(volumes, VOL_MA_LEN) if has_volume else [None] * len(closes)

    position    = 0
    eff_entry   = 0.0
    hard_stop   = 0.0
    trail_atr   = 0.0
    max_close   = 0.0
    capital     = 1.0
    trade_rets  = []
    curve       = []
    trade_log   = []
    entry_idx   = -1

    for i in range(len(closes)):
        c  = closes[i]
        lo = lows[i]
        m  = ma[i]

        if m is None or atr[i] is None:
            curve.append(capital)
            continue

        if position == 0:
            # 基础突破信号
            if i > 0 and ma[i - 1] is not None and \
               closes[i - 1] <= ma[i - 1] and c > m:
                if passes_entry_filters(i, closes, ma, ma_long, atr,
                                         volumes, vol_ma, has_volume):
                    eff_entry = c * (1 + COST)
                    hard_stop = lo
                    trail_atr = c - TRAIL_ATR_K * atr[i]
                    max_close = c
                    entry_idx = i
                    position  = 1
            curve.append(capital)
        else:
            # 持仓：先更新动态止损
            if c > max_close:
                max_close = c
            new_trail = c - TRAIL_ATR_K * atr[i]
            if new_trail > trail_atr:
                trail_atr = new_trail
            chandelier = max_close - CHANDELIER_K * atr[i]

            # 出场判定（按优先级：硬止损 → ATR 移动 → Chandelier）
            exit_reason = None
            exit_px_raw = None

            if lo <= hard_stop:
                exit_reason = "硬止损"
                exit_px_raw = hard_stop
            elif lo <= trail_atr:
                exit_reason = "ATR移动"
                exit_px_raw = trail_atr
            elif lo <= chandelier:
                exit_reason = "Chandelier"
                exit_px_raw = chandelier

            if exit_reason:
                eff_exit = exit_px_raw * (1 - COST)
                ret      = eff_exit / eff_entry - 1
                risk     = eff_entry - hard_stop
                reward   = eff_exit  - eff_entry
                rr       = (reward / risk) if (risk > 0 and ret > 0) else None
                capital *= (eff_exit / eff_entry)
                trade_rets.append(ret)
                trade_log.append((entry_idx, i, ret, exit_reason,
                                  hard_stop, exit_px_raw, rr))
                position = 0
                curve.append(capital)
            else:
                curve.append(capital * (c * (1 - COST) / eff_entry))

    # 末日强平
    if position == 1:
        exit_px_raw = closes[-1]
        eff_exit    = exit_px_raw * (1 - COST)
        ret         = eff_exit / eff_entry - 1
        risk        = eff_entry - hard_stop
        reward      = eff_exit  - eff_entry
        rr          = (reward / risk) if (risk > 0 and ret > 0) else None
        capital    *= (eff_exit / eff_entry)
        trade_rets.append(ret)
        trade_log.append((entry_idx, len(closes) - 1, ret, "强制平仓",
                          hard_stop, exit_px_raw, rr))
        curve[-1] = capital

    if not trade_rets:
        return None

    total_return = (capital - 1) * 100
    n_trades     = len(trade_rets)
    wins         = [r for r in trade_rets if r > 0]
    losses       = [r for r in trade_rets if r <= 0]
    win_rate     = len(wins) / n_trades * 100
    max_profit   = max(trade_rets) * 100
    avg_profit   = (sum(wins) / len(wins) * 100) if wins else 0.0

    sum_win  = sum(wins)
    sum_loss = abs(sum(losses))
    profit_factor = (sum_win / sum_loss) if sum_loss > 0 else None

    peak   = curve[0] if curve else 1.0
    max_dd = 0.0
    for v in curve:
        if v > peak:
            peak = v
        if peak > 0:
            dd = (peak - v) / peak * 100
            if dd > max_dd:
                max_dd = dd

    rr_values = [t[6] for t in trade_log if t[6] is not None]
    avg_rr    = (sum(rr_values) / len(rr_values)) if rr_values else None

    return {
        "n":             n,
        "total_return":  total_return,
        "n_trades":      n_trades,
        "max_drawdown":  max_dd,
        "win_rate":      win_rate,
        "max_profit":    max_profit,
        "avg_profit":    avg_profit,
        "avg_rr":        avg_rr,
        "profit_factor": profit_factor,
        "trade_log":     trade_log,
        "dates":         dates,
        "closes":        closes,
    }


def backtest_symbol(ticker):
    print(f"  下载: {ticker} ...", end=" ", flush=True)
    df = yf.download(ticker, period="max", auto_adjust=True, progress=False)
    if df is None or df.empty or len(df) < TREND_FALLBACK + ATR_LEN + 5:
        print("数据不足，跳过")
        return None, None

    closes  = df["Close"].squeeze().tolist()
    highs   = df["High"].squeeze().tolist()
    lows    = df["Low"].squeeze().tolist()
    dates   = df.index.tolist()
    if "Volume" in df.columns:
        volumes = df["Volume"].squeeze().tolist()
        # 判断成交量是否有效（大部分非零）
        non_zero = sum(1 for v in volumes if v and v > 0)
        has_volume = non_zero > len(volumes) * 0.5
    else:
        volumes = [0.0] * len(closes)
        has_volume = False

    print(f"{len(df)} 条 ({str(dates[0])[:10]}~{str(dates[-1])[:10]})"
          f"{'  无成交量' if not has_volume else ''}")

    results = []
    for n in MA_PERIODS:
        r = run_single(highs, lows, closes, volumes, dates, n, has_volume)
        if r:
            results.append(r)

    if not results:
        return None, None
    best = max(results, key=lambda x: x["total_return"])
    return results, best


# ── Excel 样式 ────────────────────────────────────────────────────────────────

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
SUBHDR_FILL = PatternFill("solid", fgColor="2E75B6")
BEST_FILL   = PatternFill("solid", fgColor="FFD966")
SEC_FILL    = PatternFill("solid", fgColor="2E75B6")
COL_FILL    = PatternFill("solid", fgColor="4472C4")
ALT_FILL    = PatternFill("solid", fgColor="F5F5F5")
WHITE_FILL  = PatternFill("solid", fgColor="FFFFFF")
WIN_FILL    = PatternFill("solid", fgColor="E2EFDA")
LOSS_FILL   = PatternFill("solid", fgColor="FCE4D6")

thin   = Side(style="thin", color="CCCCCC")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def _cell(ws, row, col, value=None, fmt=None, bold=False, size=10,
          color="000000", fill=None, align="center", wrap=False):
    cell = ws.cell(row, col)
    cell.value     = value
    cell.font      = Font(name="Arial", bold=bold, size=size, color=color)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    cell.border    = BORDER
    if fill:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    return cell


def _hdr(ws, row, col, text, fill=HEADER_FILL, color="FFFFFF", size=10):
    return _cell(ws, row, col, text, bold=True, color=color, fill=fill, size=size)


def color_ret(val):
    if val is None:
        return "000000"
    if val > 0:
        return "375623"
    if val < 0:
        return "C00000"
    return "000000"


def fmt_date(d):
    return d.strftime("%Y-%m-%d") if hasattr(d, "strftime") else str(d)[:10]


def heat_color(value, vmin, vmax):
    """红→白→绿渐变"""
    if value is None:
        return "DDDDDD"
    if vmax == vmin:
        return "FFFFFF"
    if value >= 0:
        # 0..vmax → 白(FFFFFF)..深绿(63BE7B)
        t = value / vmax if vmax > 0 else 0
        t = max(0.0, min(1.0, t))
        r = int(255 + (99  - 255) * t)
        g = int(255 + (190 - 255) * t)
        b = int(255 + (123 - 255) * t)
    else:
        t = value / vmin if vmin < 0 else 0
        t = max(0.0, min(1.0, t))
        r = int(255 + (248 - 255) * t)
        g = int(255 + (105 - 255) * t)
        b = int(255 + (107 - 255) * t)
    return f"{r:02X}{g:02X}{b:02X}"


# ── Sheet 1：汇总 ─────────────────────────────────────────────────────────────

def write_summary_sheet(wb, all_results):
    ws = wb.active
    ws.title = "汇总 - 各参数表现"

    ws.merge_cells("A1:L1")
    c = ws["A1"]
    c.value     = (f"单均线突破 v2 回测汇总  |  手续费 {FEE_RATE*100:.2f}%  "
                   f"滑点 {SLIPPAGE*100:.2f}%  |  入场 6 项过滤  |  ATR/Chandelier 出场  |  截至 {END_DATE}")
    c.font      = Font(name="Arial", bold=True, size=12, color="FFFFFF")
    c.fill      = HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    headers = ["品种", "MA N", "总收益率 %", "交易次数", "最大回撤 %",
               "胜率 %", "最大盈利 %", "平均盈利 %", "平均盈亏比 R",
               "Profit Factor", "是否最优", "备注"]
    for col, h in enumerate(headers, 1):
        _hdr(ws, 2, col, h, fill=SUBHDR_FILL)
    ws.row_dimensions[2].height = 22

    row = 3
    for sym, (results, best) in all_results.items():
        name = SYMBOLS.get(sym, sym)
        if results is None:
            _cell(ws, row, 1, f"{sym} ({name})", align="left")
            ws.merge_cells(f"B{row}:L{row}")
            _cell(ws, row, 2, "无数据", color="C00000")
            row += 1
            continue

        start_row = row
        for r in results:
            is_best = r["n"] == best["n"]
            fill    = WIN_FILL if is_best else (ALT_FILL if row % 2 == 0 else WHITE_FILL)

            _cell(ws, row, 1,  f"{sym}\n({name})", align="left", fill=fill)
            _cell(ws, row, 2,  r["n"], fmt="0", fill=fill)
            _cell(ws, row, 3,  r["total_return"], fmt='0.00"%"',
                  color=color_ret(r["total_return"]), bold=is_best, fill=fill)
            _cell(ws, row, 4,  r["n_trades"], fmt="0", fill=fill)
            _cell(ws, row, 5,  r["max_drawdown"], fmt='0.00"%"',
                  color="C00000" if r["max_drawdown"] > 50 else "000000", fill=fill)
            _cell(ws, row, 6,  r["win_rate"], fmt='0.0"%"', fill=fill)
            _cell(ws, row, 7,  r["max_profit"], fmt='0.00"%"',
                  color=color_ret(r["max_profit"]), fill=fill)
            _cell(ws, row, 8,  r["avg_profit"], fmt='0.00"%"',
                  color=color_ret(r["avg_profit"]), fill=fill)
            avg_rr = r["avg_rr"]
            _cell(ws, row, 9,
                  round(avg_rr, 2) if avg_rr is not None else "-",
                  fmt="0.00" if avg_rr is not None else None, fill=fill)
            pf = r["profit_factor"]
            _cell(ws, row, 10,
                  round(pf, 2) if pf is not None else "-",
                  fmt="0.00" if pf is not None else None,
                  color="375623" if (pf is not None and pf >= 1.5) else "000000",
                  bold=(pf is not None and pf >= 1.5), fill=fill)
            _cell(ws, row, 11, "★ 最优" if is_best else "",
                  bold=is_best, color="375623" if is_best else "000000", fill=fill)
            _cell(ws, row, 12, "", fill=fill)
            row += 1

        if row - start_row > 1:
            ws.merge_cells(f"A{start_row}:A{row-1}")
            ws["A" + str(start_row)].alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True)

    widths = [18, 8, 13, 10, 13, 10, 14, 13, 13, 13, 10, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"


# ── Sheet 2：最优排名 ────────────────────────────────────────────────────────

def write_best_sheet(wb, all_results):
    ws = wb.create_sheet("最优参数汇总")
    ws.merge_cells("A1:K1")
    c = ws["A1"]
    c.value     = "各品种最优均线参数排名（按总收益率）"
    c.font      = Font(name="Arial", bold=True, size=13, color="FFFFFF")
    c.fill      = HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    headers = ["排名", "代码", "名称", "最优 MA", "总收益率 %",
               "交易次数", "最大回撤 %", "胜率 %", "平均盈利 %",
               "平均盈亏比 R", "Profit Factor"]
    for col, h in enumerate(headers, 1):
        _hdr(ws, 2, col, h, fill=SUBHDR_FILL)

    ranked = [(sym, best) for sym, (_, best) in all_results.items() if best]
    ranked.sort(key=lambda x: x[1]["total_return"], reverse=True)

    for rank, (sym, best) in enumerate(ranked, 1):
        row  = rank + 2
        fill = ALT_FILL if rank % 2 == 0 else WHITE_FILL
        _cell(ws, row, 1, rank, fmt="0", fill=fill, bold=True)
        _cell(ws, row, 2, sym, fill=fill)
        _cell(ws, row, 3, SYMBOLS.get(sym, sym), fill=fill)
        _cell(ws, row, 4, best["n"], fmt="0", fill=fill, bold=True)
        _cell(ws, row, 5, best["total_return"], fmt='0.00"%"',
              color=color_ret(best["total_return"]), bold=True, fill=fill)
        _cell(ws, row, 6, best["n_trades"], fmt="0", fill=fill)
        _cell(ws, row, 7, best["max_drawdown"], fmt='0.00"%"',
              color="C00000" if best["max_drawdown"] > 50 else "000000", fill=fill)
        _cell(ws, row, 8, best["win_rate"], fmt='0.0"%"', fill=fill)
        _cell(ws, row, 9, best["avg_profit"], fmt='0.00"%"',
              color=color_ret(best["avg_profit"]), fill=fill)
        rr = best["avg_rr"]
        _cell(ws, row, 10,
              round(rr, 2) if rr is not None else "-",
              fmt="0.00" if rr is not None else None, fill=fill)
        pf = best["profit_factor"]
        _cell(ws, row, 11,
              round(pf, 2) if pf is not None else "-",
              fmt="0.00" if pf is not None else None,
              color="375623" if (pf is not None and pf >= 1.5) else "000000",
              bold=(pf is not None and pf >= 1.5), fill=fill)

    widths = [6, 12, 14, 10, 13, 10, 13, 10, 13, 13, 13]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"


# ── Sheet 3：稳定性热力图 ─────────────────────────────────────────────────────

def write_heatmap_sheet(wb, all_results):
    ws = wb.create_sheet("参数稳定性热力图")

    span = 2 + len(MA_PERIODS)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    c = ws["A1"]
    c.value     = "总收益率（%）参数稳定性热力图  |  绿=高 / 红=低 / 平台区即稳健最优"
    c.font      = Font(name="Arial", bold=True, size=12, color="FFFFFF")
    c.fill      = HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    _hdr(ws, 2, 1, "代码",    fill=SUBHDR_FILL)
    _hdr(ws, 2, 2, "名称",    fill=SUBHDR_FILL)
    for j, n in enumerate(MA_PERIODS, 3):
        _hdr(ws, 2, j, f"N={n}", fill=SUBHDR_FILL)

    # 收集所有数值，确定全局 vmin/vmax 用于配色
    all_vals = []
    for sym, (results, _) in all_results.items():
        if not results:
            continue
        for r in results:
            all_vals.append(r["total_return"])
    if all_vals:
        vmax = max(all_vals)
        vmin = min(all_vals)
    else:
        vmax, vmin = 1.0, -1.0

    row = 3
    for sym, (results, best) in all_results.items():
        name = SYMBOLS.get(sym, sym)
        _cell(ws, row, 1, sym,  align="center")
        _cell(ws, row, 2, name, align="left")
        by_n = {r["n"]: r for r in results} if results else {}
        for j, n in enumerate(MA_PERIODS, 3):
            r = by_n.get(n)
            if r is None:
                _cell(ws, row, j, "-", color="888888", fill=PatternFill("solid", fgColor="EEEEEE"))
                continue
            val = r["total_return"]
            color_hex = heat_color(val, vmin, vmax)
            is_best = best is not None and r["n"] == best["n"]
            _cell(ws, row, j, val,
                  fmt='0.0"%"',
                  fill=PatternFill("solid", fgColor=color_hex),
                  bold=is_best,
                  color="000000")
        row += 1

    ws.column_dimensions[get_column_letter(1)].width = 12
    ws.column_dimensions[get_column_letter(2)].width = 14
    for j in range(3, 3 + len(MA_PERIODS)):
        ws.column_dimensions[get_column_letter(j)].width = 9
    ws.freeze_panes = "C3"


# ── Sheet 4+：逐笔明细 ────────────────────────────────────────────────────────

TRADE_COLS = ["入场日期", "出场日期", "出场原因", "入场有效价", "出场有效价",
              "单笔收益率 %", "盈亏比 R", "盈亏"]
N_TRADE_COLS = len(TRADE_COLS)


def write_trade_sheet(wb, sym, results, best_n):
    safe = sym.replace("/", "-").replace(":", "-")[:28]
    ws = wb.create_sheet(safe)

    ws.merge_cells(f"A1:{get_column_letter(N_TRADE_COLS)}1")
    c = ws["A1"]
    c.value     = (f"{sym} ({SYMBOLS.get(sym, sym)})  |  手续费 {FEE_RATE*100:.2f}% "
                   f"滑点 {SLIPPAGE*100:.2f}%  |  最优 MA={best_n}")
    c.font      = Font(name="Arial", bold=True, size=12, color="FFFFFF")
    c.fill      = HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    cur = 2
    for r in results:
        n       = r["n"]
        is_best = (n == best_n)
        trades  = r["trade_log"]
        dates   = r["dates"]
        closes  = r["closes"]

        sec_fill  = BEST_FILL if is_best else SEC_FILL
        sec_color = "8B4513" if is_best else "FFFFFF"
        ws.merge_cells(f"A{cur}:{get_column_letter(N_TRADE_COLS)}{cur}")
        title = ws.cell(cur, 1)
        rr_part = f"  |  平均盈亏比 {r['avg_rr']:.2f}R" if r["avg_rr"] else ""
        pf_part = f"  |  PF {r['profit_factor']:.2f}" if r["profit_factor"] is not None else ""
        title.value = (
            f"{'★ 最优  ' if is_best else ''}MA = {n}  |  "
            f"总收益 {r['total_return']:.2f}%  |  "
            f"交易 {r['n_trades']} 笔  |  "
            f"胜率 {r['win_rate']:.1f}%  |  "
            f"最大回撤 {r['max_drawdown']:.1f}%"
            f"{rr_part}{pf_part}"
        )
        title.font      = Font(name="Arial", bold=True, size=10, color=sec_color)
        title.fill      = sec_fill
        title.alignment = Alignment(horizontal="left", vertical="center")
        title.border    = BORDER
        ws.row_dimensions[cur].height = 20
        cur += 1

        for col, h in enumerate(TRADE_COLS, 1):
            _hdr(ws, cur, col, h, fill=COL_FILL, size=9)
        ws.row_dimensions[cur].height = 18
        cur += 1

        if not trades:
            ws.merge_cells(f"A{cur}:{get_column_letter(N_TRADE_COLS)}{cur}")
            _cell(ws, cur, 1, "该参数下无交易记录", color="888888",
                  fill=WHITE_FILL, align="left")
            cur += 1
        else:
            for (ei, xi, ret, reason, hard_stop, exit_px_raw, rr) in trades:
                fill = WIN_FILL if ret > 0 else LOSS_FILL
                entry_date = fmt_date(dates[ei])
                exit_date  = fmt_date(dates[xi])
                entry_px   = closes[ei] * (1 + COST)
                exit_px    = exit_px_raw * (1 - COST)

                reason_color = "C00000" if reason in ("硬止损",) else (
                    "1F4E79" if reason in ("ATR移动", "Chandelier") else "000000")

                _cell(ws, cur, 1, entry_date, fill=fill)
                _cell(ws, cur, 2, exit_date, fill=fill)
                _cell(ws, cur, 3, reason, fill=fill, color=reason_color)
                _cell(ws, cur, 4, round(entry_px, 6), fmt="#,##0.######", fill=fill)
                _cell(ws, cur, 5, round(exit_px, 6),  fmt="#,##0.######", fill=fill)
                _cell(ws, cur, 6, ret * 100, fmt='0.00"%"',
                      color=color_ret(ret), bold=(ret > 0), fill=fill)
                if rr is not None:
                    _cell(ws, cur, 7, round(rr, 2), fmt="0.00",
                          color="375623", bold=True, fill=fill)
                else:
                    _cell(ws, cur, 7, "-", color="888888", fill=fill)
                _cell(ws, cur, 8, "盈" if ret > 0 else "亏",
                      color="375623" if ret > 0 else "C00000",
                      bold=True, fill=fill)
                cur += 1

        cur += 1

    widths = [14, 14, 12, 16, 16, 15, 12, 8]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


# ── 主流程 ────────────────────────────────────────────────────────────────────

def main():
    all_results = {}
    for sym in SYMBOLS:
        print(f"\n[{sym}]")
        results, best = backtest_symbol(sym)
        all_results[sym] = (results, best)
        if best:
            rr_str = f"  R={best['avg_rr']:.2f}" if best["avg_rr"] else ""
            pf_str = f"  PF={best['profit_factor']:.2f}" if best["profit_factor"] is not None else ""
            print(f"  最优 MA={best['n']}  收益={best['total_return']:.1f}%  "
                  f"胜率={best['win_rate']:.1f}%  回撤={best['max_drawdown']:.1f}%{rr_str}{pf_str}")

    out_path = "/Users/zhangzheyu/market-replay/data/ma_breakout_v2_backtest.xlsx"
    os.makedirs(os.path.dirname(out_path), exist_ok=True)

    print("\n生成 Excel ...")
    wb = openpyxl.Workbook()
    write_summary_sheet(wb, all_results)
    write_best_sheet(wb, all_results)
    write_heatmap_sheet(wb, all_results)
    for sym, (results, best) in all_results.items():
        if results and best:
            write_trade_sheet(wb, sym, results, best["n"])

    wb.save(out_path)
    print(f"✅ 已保存至 {out_path}")


if __name__ == "__main__":
    main()
