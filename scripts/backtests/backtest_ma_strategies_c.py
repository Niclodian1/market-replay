"""
策略 C：多均线同时上穿（C-x2 / C-x3 / C-x4）

入场（同一根 K 线收盘时）：
  - 均线池 {30, 60, 99, 120}，对每根 MAi 判定上穿：prev_c ≤ prev_MAi 且 c > MAi
  - crossed = 本根满足上穿的均线集合，k = len(crossed)
  - 实体涨幅 c/open − 1 > 3%
  - 恰好 k 根上穿 → 触发变体 C-x{k}（k∈{2,3,4}，三变体互斥）
  - MA_exit = 被上穿的 N 根中周期最小那根（每笔仓位各自不同）

出场（与策略 A 一致，按优先级）：
  1. low ≤ 止损价 → 止损 / 保本
  2. 收盘低于该笔 MA_exit 的幅度 ≥ 3% → 均线出场（可累积）
  止损：初始 = 入场 K 线 low；浮盈 > 4% 后下一根移至成本价（保本）。

数据：BTC-USD / ETH-USD × 1day / 4h / 1h，经 LocalCryptoProvider（parquet + OKX 增量）。
产出：backend/data/ma_strategies_c/{SYMBOL}_{tf}.json + _manifest.json + Excel 逐笔明细。
"""

import argparse
import asyncio
import datetime
import os
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

# ── 让脚本能 import backend app（参考 backtest_w_macd.py） ──────────────────────
ROOT = Path(__file__).resolve().parents[2]  # repo root
sys.path.insert(0, str(ROOT / "backend"))
_SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(_SCRIPT_DIR))

from app.providers.local_crypto import LocalCryptoProvider  # noqa: E402

# ── 复用策略 A/B 的共享工具与 Excel 样式 ─────────────────────────────────────
from backtest_ma_strategies import (  # noqa: E402
    COST, FEE_RATE, SLIPPAGE, BODY_RISE_MIN, BREAKEVEN_TRIGGER,
    rolling_sma, summarize,
    _cell, _hdr, color_ret,
    HEADER_FILL, SUBHDR_FILL, BEST_FILL, SEC_FILL, COL_FILL,
    ALT_FILL, WHITE_FILL, WIN_FILL, LOSS_FILL,
)
from backtest_json_export import _to_unix, _build_candles, _build_line  # noqa: E402

# ── 配置 ─────────────────────────────────────────────────────────────────────
SYMBOLS = {"BTC-USD": "BTC", "ETH-USD": "ETH"}
TIMEFRAMES = ["1day", "4h", "1h"]
MA_POOL = [30, 60, 99, 120]
VARIANTS = [2, 3, 4]            # C-x2 / C-x3 / C-x4
MA_EXIT_PCT = 0.03
END_DATE = datetime.date.today().isoformat()

JSON_DIR = ROOT / "backend" / "data" / "ma_strategies_c"
XLSX_PATH = _SCRIPT_DIR / "data" / "ma_strategies_c_backtest.xlsx"


def variant_label(k):
    return f"x{k}"


# ── 数据加载 ─────────────────────────────────────────────────────────────────

def load_candles(symbol, tf):
    """经 LocalCryptoProvider 取全量（parquet + OKX 增量）。返回时间升序的 OHLC 列表。"""
    candles = asyncio.run(LocalCryptoProvider().get_full_history(symbol, tf))
    return candles


# ── 出场判定（单笔仓位，MA_exit 为该笔自有线） ──────────────────────────────

def _try_exit_c(p, c, lo, ma_exit_val):
    """更新浮动止损/最高价；触发出场返回 (eff_exit, reason)，否则 (None, None)。
    逻辑与策略 A _try_exit 一致，仅 ma 换成该笔 MA_exit 在当根的值。"""
    if p.get("breakeven_pending"):
        p["stop_loss"] = p["entry_close"]
        p["breakeven_set"] = True
        p["breakeven_pending"] = False

    if c > p["max_close"]:
        p["max_close"] = c
    if not p["breakeven_set"] and not p.get("breakeven_pending") and c / p["entry_close"] - 1 > BREAKEVEN_TRIGGER:
        p["breakeven_pending"] = True

    ma_exit = (ma_exit_val is not None and ma_exit_val > 0 and (ma_exit_val - c) / ma_exit_val >= MA_EXIT_PCT)
    stop_exit = lo <= p["stop_loss"]

    if stop_exit:
        return p["stop_loss"] * (1 - COST), ("保本" if p["breakeven_set"] else "止损")
    if ma_exit:
        return c * (1 - COST), "均线"
    return None, None


def _close_position_c(p, i, eff_exit, reason, trade_rets, trade_log):
    ret = eff_exit / p["eff_entry"] - 1
    risk = p["eff_entry"] - p["initial_stop_loss"]
    reward = eff_exit - p["eff_entry"]
    rr = (reward / risk) if (risk > 0 and ret > 0) else None
    max_ret = p["max_close"] * (1 - COST) / p["eff_entry"] - 1
    trade_rets.append(ret)
    # 元组比 A/B 多两位：exit_ma_period、crossed（升序元组）
    trade_log.append((p["entry_idx"], i, ret, reason, p["stop_loss"], rr, max_ret,
                      p["exit_ma_period"], tuple(p["crossed"])))
    return ret


# ── 单变体回测（并行持仓） ───────────────────────────────────────────────────

def run_variant(closes, opens, lows, dates, sma_by_n, crossed_by_bar, target_k):
    open_positions = []
    trade_rets, trade_log, curve = [], [], []
    realized = 0.0
    n = len(closes)

    for i in range(n):
        c, lo = closes[i], lows[i]

        # 1) 先处理已有持仓的出场（各自看自己的 MA_exit）
        still = []
        for p in open_positions:
            ma_exit_val = sma_by_n[p["exit_ma_period"]][i]
            eff_exit, reason = _try_exit_c(p, c, lo, ma_exit_val)
            if reason:
                realized += _close_position_c(p, i, eff_exit, reason, trade_rets, trade_log)
            else:
                still.append(p)
        open_positions = still

        # 2) 再判断开新仓：恰好 target_k 根上穿 + 实体涨幅 > 3%
        crossed = crossed_by_bar[i]
        if len(crossed) == target_k:
            op = opens[i]
            body_rise = (c / op - 1) if op and op > 0 else 0
            if body_rise > BODY_RISE_MIN:
                open_positions.append({
                    "eff_entry": c * (1 + COST), "entry_close": c,
                    "stop_loss": lo, "initial_stop_loss": lo,
                    "breakeven_set": False, "breakeven_pending": False,
                    "max_close": c, "entry_idx": i,
                    "exit_ma_period": min(crossed), "crossed": sorted(crossed),
                })

        unreal = sum((c * (1 - COST) / p["eff_entry"] - 1) for p in open_positions)
        curve.append(1.0 + realized + unreal)

    # 末尾强制平仓
    last_i = n - 1
    for p in open_positions:
        eff_exit = closes[-1] * (1 - COST)
        realized += _close_position_c(p, last_i, eff_exit, "强制平仓", trade_rets, trade_log)
    if curve:
        curve[-1] = 1.0 + realized

    return summarize(target_k, realized, trade_rets, trade_log, curve, dates, closes)


def _compute_crossed_by_bar(closes, sma_by_n):
    """每根 K 线上穿的均线集合（来自 MA_POOL）。"""
    n = len(closes)
    out = [[] for _ in range(n)]
    for i in range(1, n):
        c, pc = closes[i], closes[i - 1]
        crossed = []
        for m in MA_POOL:
            ma, pma = sma_by_n[m][i], sma_by_n[m][i - 1]
            if ma is not None and pma is not None and pc <= pma and c > ma:
                crossed.append(m)
        out[i] = crossed
    return out


def backtest_symbol_tf(symbol, tf):
    candles = load_candles(symbol, tf)
    if not candles or len(candles) < max(MA_POOL) + 10:
        print(f"  [{symbol} {tf}] 数据不足，跳过")
        return None

    times = [c["time"] for c in candles]
    opens = [c["open"] for c in candles]
    highs = [c["high"] for c in candles]
    lows = [c["low"] for c in candles]
    closes = [c["close"] for c in candles]

    sma_by_n = {m: rolling_sma(closes, m) for m in MA_POOL}
    crossed_by_bar = _compute_crossed_by_bar(closes, sma_by_n)

    results = {}
    for k in VARIANTS:
        r = run_variant(closes, opens, lows, times, sma_by_n, crossed_by_bar, k)
        results[k] = r

    best_k = None
    valid = [(k, r) for k, r in results.items() if r]
    if valid:
        best_k = max(valid, key=lambda kr: kr[1]["total_return"])[0]

    ohlc = {
        "dates": times, "opens": opens, "highs": highs, "lows": lows, "closes": closes,
        "sma_by_n": sma_by_n,
    }
    return {"results": results, "best_k": best_k, "ohlc": ohlc}


# ── JSON 导出 ────────────────────────────────────────────────────────────────

def _build_variant_block(r, ohlc, is_best):
    dates = ohlc["dates"]
    closes = ohlc["closes"]
    trades_out = []
    for ei, xi, ret, reason, sl, rr, max_ret, exit_ma, crossed in r["trade_log"]:
        entry_px = closes[ei] * (1 + COST)
        if reason in ("止损", "保本"):
            exit_px = sl * (1 - COST)
        else:
            exit_px = closes[xi] * (1 - COST)
        trades_out.append({
            "entry_time": _to_unix(dates[ei]),
            "exit_time": _to_unix(dates[xi]),
            "entry_price": round(entry_px, 6),
            "exit_price": round(exit_px, 6),
            "stop_loss": round(sl, 6),
            "pnl_pct": round(ret * 100, 4),
            "max_profit_pct": round(max_ret * 100, 4),
            "rr": round(rr, 4) if rr is not None else None,
            "exit_reason": reason,
            "exit_ma": exit_ma,
            "crossed": list(crossed),
        })
    return {
        "stats": {
            "total_return": round(r["total_return"], 4),
            "n_trades": r["n_trades"],
            "win_rate": round(r["win_rate"], 4),
            "max_drawdown": round(r["max_drawdown"], 4),
            "max_profit": round(r["max_profit"], 4),
            "avg_profit": round(r["avg_profit"], 4),
            "avg_rr": round(r["avg_rr"], 4) if r["avg_rr"] is not None else None,
            "is_best": is_best,
        },
        "trades": trades_out,
    }


def export_symbol_tf_json(symbol, tf, data, out_path):
    import json
    ohlc = data["ohlc"]
    best_k = data["best_k"]
    variants = {}
    for k in VARIANTS:
        r = data["results"].get(k)
        if r:
            variants[variant_label(k)] = _build_variant_block(r, ohlc, is_best=(k == best_k))
        else:
            variants[variant_label(k)] = {
                "stats": {"total_return": 0, "n_trades": 0, "win_rate": 0, "max_drawdown": 0,
                          "max_profit": 0, "avg_profit": 0, "avg_rr": None, "is_best": False},
                "trades": [],
            }
    payload = {
        "symbol": symbol,
        "name": SYMBOLS.get(symbol, symbol),
        "timeframe": tf,
        "generated_at": datetime.datetime.now().isoformat(timespec="seconds"),
        "candles": _build_candles(ohlc),
        "ma_lines": {str(m): _build_line(line, ohlc["dates"]) for m, line in ohlc["sma_by_n"].items()},
        "variants": variants,
    }
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, separators=(",", ":"))


def write_manifest():
    import json
    payload = {
        "generated_at": datetime.datetime.now().isoformat(timespec="seconds"),
        "symbols": [{"code": code, "name": name} for code, name in SYMBOLS.items()],
        "timeframes": TIMEFRAMES,
        "variants": [variant_label(k) for k in VARIANTS],
    }
    with open(JSON_DIR / "_manifest.json", "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)


# ── Excel ────────────────────────────────────────────────────────────────────

def _fmt_dt(unix, tf):
    d = datetime.datetime.utcfromtimestamp(int(unix))
    return d.strftime("%Y-%m-%d") if tf == "1day" else d.strftime("%Y-%m-%d %H:%M")


def write_summary_sheet_c(wb, all_results):
    ws = wb.create_sheet("策略C 汇总")
    ws.merge_cells("A1:K1")
    c = ws["A1"]
    c.value = (f"策略C 多均线同时上穿  |  手续费 {FEE_RATE*100:.2f}%  滑点 {SLIPPAGE*100:.2f}%  "
               f"均线池 {MA_POOL}  实体涨幅>{BODY_RISE_MIN*100:.0f}%  均线出场≥{MA_EXIT_PCT*100:.0f}%  |  截至 {END_DATE}")
    c.font = Font(name="Arial", bold=True, size=12, color="FFFFFF")
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    headers = ["品种", "周期", "变体", "总收益率 %", "交易次数", "最大回撤 %",
               "胜率 %", "最大单笔盈利 %", "平均盈利 %", "平均盈亏比 R", "是否最优"]
    for col, h in enumerate(headers, 1):
        _hdr(ws, 2, col, h, fill=SUBHDR_FILL)
    ws.row_dimensions[2].height = 22

    row = 3
    for (symbol, tf), data in all_results.items():
        if not data:
            continue
        name = SYMBOLS.get(symbol, symbol)
        best_k = data["best_k"]
        start_row = row
        for k in VARIANTS:
            r = data["results"].get(k)
            is_best = (k == best_k)
            fill = WIN_FILL if is_best else (ALT_FILL if row % 2 == 0 else WHITE_FILL)
            _cell(ws, row, 1, f"{symbol}\n({name})", align="left", fill=fill)
            _cell(ws, row, 2, tf, fill=fill)
            _cell(ws, row, 3, f"C-{variant_label(k)}", bold=is_best, fill=fill)
            if r:
                _cell(ws, row, 4, r["total_return"], fmt='0.00"%"',
                      color=color_ret(r["total_return"]), bold=is_best, fill=fill)
                _cell(ws, row, 5, r["n_trades"], fmt="0", fill=fill)
                _cell(ws, row, 6, r["max_drawdown"], fmt='0.00"%"',
                      color="C00000" if r["max_drawdown"] > 50 else "000000", fill=fill)
                _cell(ws, row, 7, r["win_rate"], fmt='0.0"%"', fill=fill)
                _cell(ws, row, 8, r["max_profit"], fmt='0.00"%"', color=color_ret(r["max_profit"]), fill=fill)
                _cell(ws, row, 9, r["avg_profit"], fmt='0.00"%"', color=color_ret(r["avg_profit"]), fill=fill)
                avg_rr = r["avg_rr"]
                _cell(ws, row, 10, round(avg_rr, 2) if avg_rr is not None else "-",
                      fmt="0.00" if avg_rr is not None else None, fill=fill)
            else:
                for col in range(4, 11):
                    _cell(ws, row, col, "-", fill=fill)
            _cell(ws, row, 11, "★ 最优" if is_best else "",
                  bold=is_best, color="375623" if is_best else "000000", fill=fill)
            row += 1
        if row - start_row > 1:
            ws.merge_cells(f"A{start_row}:A{row-1}")
            ws[f"A{start_row}"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    widths = [16, 8, 9, 13, 10, 13, 10, 16, 13, 13, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"


TRADE_COLS_C = ["入场日期", "出场日期", "出场原因", "入场有效价", "出场有效价",
                "单笔收益率 %", "最大盈利率 %", "盈亏比 R", "出场参考线 MA", "上穿均线", "盈亏"]
N_TRADE_COLS_C = len(TRADE_COLS_C)


def write_trade_sheet_c(wb, symbol, tf, data):
    name = SYMBOLS.get(symbol, symbol)
    safe = f"C-{name}-{tf}".replace("/", "-")[:28]
    ws = wb.create_sheet(safe)

    ws.merge_cells(f"A1:{get_column_letter(N_TRADE_COLS_C)}1")
    c = ws["A1"]
    c.value = (f"策略C | {symbol} ({name}) · {tf}  —  手续费 {FEE_RATE*100:.2f}%  滑点 {SLIPPAGE*100:.2f}%  "
               f"均线池 {MA_POOL}")
    c.font = Font(name="Arial", bold=True, size=12, color="FFFFFF")
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    ohlc = data["ohlc"]
    dates = ohlc["dates"]
    closes = ohlc["closes"]
    best_k = data["best_k"]

    cur_row = 2
    for k in VARIANTS:
        r = data["results"].get(k)
        is_best = (k == best_k)
        sec_fill = BEST_FILL if is_best else SEC_FILL
        sec_color = "8B4513" if is_best else "FFFFFF"
        ws.merge_cells(f"A{cur_row}:{get_column_letter(N_TRADE_COLS_C)}{cur_row}")
        title_cell = ws.cell(cur_row, 1)
        _prefix = f"{'★ 最优  ' if is_best else ''}C-{variant_label(k)}  |  "
        if r:
            _base = (f"总收益 {r['total_return']:.2f}%  |  交易 {r['n_trades']} 笔  |  "
                     f"胜率 {r['win_rate']:.1f}%  |  最大回撤 {r['max_drawdown']:.1f}%  |  "
                     f"最大单笔盈利 {r['max_profit']:.2f}%")
            if r.get("avg_rr"):
                _base += f"  |  平均盈亏比 {r['avg_rr']:.2f}R"
        else:
            _base = "无有效交易"
        title_cell.value = _prefix + _base
        title_cell.font = Font(name="Arial", bold=True, size=10, color=sec_color)
        title_cell.fill = sec_fill
        title_cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[cur_row].height = 20
        cur_row += 1

        for col, h in enumerate(TRADE_COLS_C, 1):
            _hdr(ws, cur_row, col, h, fill=COL_FILL, size=9)
        ws.row_dimensions[cur_row].height = 18
        cur_row += 1

        trades = r["trade_log"] if r else []
        if not trades:
            ws.merge_cells(f"A{cur_row}:{get_column_letter(N_TRADE_COLS_C)}{cur_row}")
            _cell(ws, cur_row, 1, "该变体下无交易记录", color="888888", fill=WHITE_FILL, align="left")
            cur_row += 1
        else:
            for (ei, xi, ret, reason, sl, rr, max_ret, exit_ma, crossed) in trades:
                fill = WIN_FILL if ret > 0 else LOSS_FILL
                entry_px = closes[ei] * (1 + COST)
                exit_px = sl * (1 - COST) if reason in ("止损", "保本") else closes[xi] * (1 - COST)
                reason_color = ("C00000" if reason == "止损"
                                else "D06000" if reason == "保本" else "000000")
                _cell(ws, cur_row, 1, _fmt_dt(dates[ei], tf), fill=fill)
                _cell(ws, cur_row, 2, _fmt_dt(dates[xi], tf), fill=fill)
                _cell(ws, cur_row, 3, reason, fill=fill, color=reason_color)
                _cell(ws, cur_row, 4, round(entry_px, 6), fmt="#,##0.######", fill=fill)
                _cell(ws, cur_row, 5, round(exit_px, 6), fmt="#,##0.######", fill=fill)
                _cell(ws, cur_row, 6, ret * 100, fmt='0.00"%"', color=color_ret(ret), bold=(ret > 0), fill=fill)
                _cell(ws, cur_row, 7, max_ret * 100, fmt='0.00"%"', color="375623", bold=True, fill=fill)
                if rr is not None:
                    _cell(ws, cur_row, 8, round(rr, 2), fmt="0.00", color="375623", bold=True, fill=fill)
                else:
                    _cell(ws, cur_row, 8, "-", color="888888", fill=fill)
                _cell(ws, cur_row, 9, f"MA{exit_ma}", fill=fill)
                _cell(ws, cur_row, 10, "/".join(f"MA{m}" for m in crossed), fill=fill)
                _cell(ws, cur_row, 11, "盈" if ret > 0 else "亏",
                      color="375623" if ret > 0 else "C00000", bold=True, fill=fill)
                cur_row += 1
        cur_row += 1

    widths = [18, 18, 10, 16, 16, 14, 14, 10, 14, 18, 8]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


# ── 主流程 ───────────────────────────────────────────────────────────────────

def main(only_symbol=None, only_tf=None, skip_excel=False):
    if only_symbol and only_symbol not in SYMBOLS:
        raise SystemExit(f"未知品种：{only_symbol}（合法范围：{', '.join(SYMBOLS)})")
    if only_tf and only_tf not in TIMEFRAMES:
        raise SystemExit(f"未知周期：{only_tf}（合法范围：{', '.join(TIMEFRAMES)})")

    symbols = [only_symbol] if only_symbol else list(SYMBOLS.keys())
    tfs = [only_tf] if only_tf else TIMEFRAMES

    os.makedirs(JSON_DIR, exist_ok=True)
    all_results = {}

    for symbol in symbols:
        print(f"\n[{symbol}]")
        for tf in tfs:
            print(f"  下载 + 回测 {tf} ...", flush=True)
            data = backtest_symbol_tf(symbol, tf)
            all_results[(symbol, tf)] = data
            if not data:
                continue
            safe = symbol.replace("/", "-")
            export_symbol_tf_json(symbol, tf, data, JSON_DIR / f"{safe}_{tf}.json")
            for k in VARIANTS:
                r = data["results"].get(k)
                star = " ★" if k == data["best_k"] else ""
                if r:
                    rr = f"  盈亏比={r['avg_rr']:.2f}R" if r["avg_rr"] else ""
                    print(f"    C-x{k}{star}: 收益={r['total_return']:8.1f}%  "
                          f"交易={r['n_trades']:4d}  胜率={r['win_rate']:5.1f}%  "
                          f"回撤={r['max_drawdown']:5.1f}%{rr}")
                else:
                    print(f"    C-x{k}{star}: 无有效交易")

    # manifest 仅全量回测时刷新
    if not only_symbol and not only_tf:
        write_manifest()
        print(f"\n✅ 已导出 JSON + manifest 到 {JSON_DIR}/")
    else:
        print(f"\n✅ 已导出 JSON 到 {JSON_DIR}/")

    # Excel 仅全量回测时生成
    if not skip_excel and not only_symbol and not only_tf:
        os.makedirs(XLSX_PATH.parent, exist_ok=True)
        print("\n生成 Excel ...")
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        write_summary_sheet_c(wb, all_results)
        for (symbol, tf), data in all_results.items():
            if data:
                write_trade_sheet_c(wb, symbol, tf, data)
        wb.save(str(XLSX_PATH))
        print(f"✅ 已保存 Excel 至 {XLSX_PATH}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--symbol", help="只回测该品种（默认 BTC/ETH 全部）")
    parser.add_argument("--timeframe", help="只回测该周期（1day/4h/1h，默认全部）")
    parser.add_argument("--no-excel", action="store_true", help="跳过 Excel 输出")
    args = parser.parse_args()
    main(only_symbol=args.symbol, only_tf=args.timeframe, skip_excel=args.no_excel)
