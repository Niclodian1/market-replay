"""
杠杆浮盈加仓策略回测 (2x / 3x)
策略 A：单均线突破，保本触发改为首次入场价 +15%
策略 B：MA120 之上均线回调，保本触发改为首次入场价 +15%
加仓规则：全仓模式——下一根K线收盘，将所有新增浮盈 × 杠杆倍数投入
Excel 对比：1x（原策略保本 4%）vs 2x vs 3x
"""

import datetime
import os
import sys
import asyncio
from pathlib import Path

import yfinance as yf
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# ── 从原策略文件导入共享工具 ─────────────────────────────────────────────────
_REPO = Path(__file__).resolve().parent
sys.path.insert(0, str(_REPO))
from backtest_ma_strategies import (
    SYMBOLS, MA_PERIODS, TREND_MA, FEE_RATE, SLIPPAGE, COST, BODY_RISE_MIN,
    rolling_sma, rolling_rsi,
    run_breakout, run_pullback,
    summarize,
    CRYPTO_SYMBOLS,
)

# ── 常量 ─────────────────────────────────────────────────────────────────────
LEVERAGE_LEVELS     = [2, 3]
BREAKEVEN_TRIGGER_3X = 0.15   # 首次入场价上涨 15% 才保本（原为 4%）
MA_EXIT_PCT         = 0.03
END_DATE            = datetime.date.today().isoformat()


# ── 杠杆仓位工具函数 ─────────────────────────────────────────────────────────

def _unrealized_pnl(sub_positions, c):
    """当前收盘价下的浮动盈亏（含出场摩擦成本），相对初始 1.0 本金。"""
    eff = c * (1 - COST)
    return sum(sp["size"] * (eff / sp["entry_price"] - 1) for sp in sub_positions)


def _exit_pnl(sub_positions, eff_exit):
    """以有效出场价（已含摩擦成本）结算的实现盈亏。"""
    return sum(sp["size"] * (eff_exit / sp["entry_price"] - 1) for sp in sub_positions)


def _breakeven_price(sub_positions):
    """让整笔交易 PnL=0 的出场价：按 size 加权的调和平均成本（已含入场摩擦）。
    数学上 Σ size_i × (P / entry_i − 1) = 0  =>  P = Σ size_i / Σ(size_i / entry_i)
    """
    total_size = sum(sp["size"] for sp in sub_positions)
    denom      = sum(sp["size"] / sp["entry_price"] for sp in sub_positions)
    if total_size <= 0 or denom <= 0:
        return 0.0
    return total_size / denom


def _new_trade(i, c, lo, leverage):
    entry_price = c * (1 + COST)
    return {
        "entry_idx":           i,
        "entry_close":         c,          # 首次入场收盘价（保本触发基准）
        "stop_loss":           lo,
        "initial_stop_loss":   lo,
        "breakeven_set":       False,
        "breakeven_pending":   False,
        "max_close":           c,
        "max_pnl":             0.0,        # 持仓期间最大浮盈
        "leverage":            leverage,
        "sub_positions": [{
            "entry_price": entry_price,
            "size":        float(leverage),  # size = margin × leverage
            "margin":      1.0,
            "bar_idx":     i,
        }],
        "cumulative_pnl_added": 0.0,       # 已安排加仓消耗的浮盈高水位
        "pending_add_margin":   0.0,       # 下一根K线要执行的加仓保证金
        "events": [{
            "bar_idx": i, "type": "entry",
            "price":   entry_price,
            "size":    float(leverage),
        }],
    }


def _process_trade(trade, i, c, lo, ma, leverage, trade_rets, trade_log):
    """
    处理单笔持仓的一根K线：返回 True 表示该仓位已平仓（不再持有），
    False 表示继续持有并已完成本棒的加仓/浮盈更新。
    """
    # ── 1. 激活上一棒触发的保本条件 ──────────────────────────────────────────
    if trade["breakeven_pending"]:
        trade["stop_loss"]       = _breakeven_price(trade["sub_positions"])
        trade["breakeven_set"]   = True
        trade["breakeven_pending"] = False

    # ── 2. 更新最高收盘价与最大浮盈 ─────────────────────────────────────────
    if c > trade["max_close"]:
        trade["max_close"] = c
    cur_pnl = _unrealized_pnl(trade["sub_positions"], c)
    if cur_pnl > trade["max_pnl"]:
        trade["max_pnl"] = cur_pnl

    # ── 3. 检查出场条件（止损用低点，均线出场用收盘）───────────────────────
    stop_hit = lo <= trade["stop_loss"]
    ma_exit  = (ma is not None and ma > 0 and (ma - c) / ma >= MA_EXIT_PCT)

    if stop_hit or ma_exit:
        if stop_hit:
            eff_exit = trade["stop_loss"] * (1 - COST)
            reason   = "保本" if trade["breakeven_set"] else "止损"
        else:
            eff_exit = c * (1 - COST)
            reason   = "均线"

        total_pnl = _exit_pnl(trade["sub_positions"], eff_exit)
        max_ret   = trade["max_pnl"]

        # 计算盈亏比（仅赢单才有意义）
        risk   = trade["sub_positions"][0]["entry_price"] - trade["initial_stop_loss"]
        reward = eff_exit - trade["sub_positions"][0]["entry_price"]
        rr     = (reward / risk) if (risk > 0 and total_pnl > 0) else None

        # 先追加出场事件，再写入 log
        trade["events"].append({
            "bar_idx": i, "type": "exit",
            "price":   eff_exit, "size": None, "reason": reason,
        })
        trade_rets.append(total_pnl)
        trade_log.append((
            trade["entry_idx"], i, total_pnl, reason,
            trade["stop_loss"], rr, max_ret,
            leverage, list(trade["events"]),
        ))
        return True   # 已平仓

    # ── 4. 执行上一棒安排的加仓（出场检查通过后才执行）────────────────────
    if trade["pending_add_margin"] > 1e-9:
        m          = trade["pending_add_margin"]
        add_size   = m * leverage
        add_price  = c * (1 + COST)
        trade["sub_positions"].append({
            "entry_price": add_price,
            "size":        add_size,
            "margin":      m,
            "bar_idx":     i,
        })
        trade["events"].append({
            "bar_idx": i, "type": "add",
            "price":   add_price, "size": add_size,
        })
        trade["pending_add_margin"] = 0.0
        if trade["breakeven_set"]:
            trade["stop_loss"] = _breakeven_price(trade["sub_positions"])

    # ── 5. 检查本棒是否触发保本条件（下一棒生效）───────────────────────────
    if (not trade["breakeven_set"] and not trade["breakeven_pending"]
            and c / trade["entry_close"] - 1 > BREAKEVEN_TRIGGER_3X):
        trade["breakeven_pending"] = True

    # ── 6. 计算浮盈增量，安排下一棒加仓 ──────────────────────────────────
    cur_pnl   = _unrealized_pnl(trade["sub_positions"], c)   # 含出场成本的估值
    increment = cur_pnl - trade["cumulative_pnl_added"]
    if increment > 1e-9:
        trade["pending_add_margin"]    = increment
        trade["cumulative_pnl_added"]  = cur_pnl

    return False  # 继续持有


def _force_close(open_trades, closes, trade_rets, trade_log, leverage):
    realized = 0.0
    last_i   = len(closes) - 1
    for trade in open_trades:
        eff_exit = closes[-1] * (1 - COST)
        total_pnl = _exit_pnl(trade["sub_positions"], eff_exit)
        trade["events"].append({
            "bar_idx": last_i, "type": "exit",
            "price": eff_exit, "size": None, "reason": "强制平仓",
        })
        trade_rets.append(total_pnl)
        trade_log.append((
            trade["entry_idx"], last_i, total_pnl, "强制平仓",
            trade["stop_loss"], None, trade["max_pnl"],
            leverage, list(trade["events"]),
        ))
        realized += total_pnl
    return realized


def summarize_3x(n, leverage, realized, trade_rets, trade_log, curve, dates, closes):
    r = summarize(n, realized, trade_rets, trade_log, curve, dates, closes)
    if r is not None:
        r["leverage"] = leverage
    return r


# ── 策略 A 杠杆版：单均线突破 ─────────────────────────────────────────────

def run_leveraged_breakout(closes, lows, opens, dates, n, leverage):
    if len(closes) <= n:
        return None
    sma = rolling_sma(closes, n)

    open_trades = []
    trade_rets, trade_log, curve = [], [], []
    realized = 0.0

    for i in range(len(closes)):
        c, lo = closes[i], lows[i]
        ma = sma[i]
        if ma is None:
            curve.append(1.0 + realized)
            continue
        prev_ma = sma[i - 1] if i > 0 else None
        prev_c  = closes[i - 1] if i > 0 else None

        # 处理现有仓位
        still_open = []
        for trade in open_trades:
            closed = _process_trade(trade, i, c, lo, ma, leverage, trade_rets, trade_log)
            if closed:
                realized += trade_log[-1][2]   # trade_log[-1][2] = total_pnl
            else:
                still_open.append(trade)
        open_trades = still_open

        # 入场信号（同原策略 A）
        if prev_ma is not None and prev_c is not None:
            op        = opens[i]
            body_rise = (c / op - 1) if op and op > 0 else 0
            if prev_c <= prev_ma and c > ma and body_rise > BODY_RISE_MIN:
                open_trades.append(_new_trade(i, c, lo, leverage))

        unreal = sum(_unrealized_pnl(t["sub_positions"], c) for t in open_trades)
        curve.append(1.0 + realized + unreal)

    realized += _force_close(open_trades, closes, trade_rets, trade_log, leverage)
    if curve:
        curve[-1] = 1.0 + realized

    return summarize_3x(n, leverage, realized, trade_rets, trade_log, curve, dates, closes)


# ── 策略 B 杠杆版：MA120 之上均线回调 ─────────────────────────────────────

def run_leveraged_pullback(closes, lows, opens, dates, n, trend_sma, rsi_vals, leverage):
    if len(closes) <= max(n, TREND_MA):
        return None
    sma = rolling_sma(closes, n)

    open_trades = []
    trade_rets, trade_log, curve = [], [], []
    realized = 0.0

    for i in range(len(closes)):
        c, lo = closes[i], lows[i]
        ma    = sma[i]
        tma   = trend_sma[i]
        rsi   = rsi_vals[i]
        if ma is None or tma is None or rsi is None:
            curve.append(1.0 + realized)
            continue

        still_open = []
        for trade in open_trades:
            closed = _process_trade(trade, i, c, lo, ma, leverage, trade_rets, trade_log)
            if closed:
                realized += trade_log[-1][2]
            else:
                still_open.append(trade)
        open_trades = still_open

        # 入场信号（同原策略 B）
        if lo <= ma and c > ma and c > tma and rsi < 50:
            open_trades.append(_new_trade(i, c, lo, leverage))

        unreal = sum(_unrealized_pnl(t["sub_positions"], c) for t in open_trades)
        curve.append(1.0 + realized + unreal)

    realized += _force_close(open_trades, closes, trade_rets, trade_log, leverage)
    if curve:
        curve[-1] = 1.0 + realized

    return summarize_3x(n, leverage, realized, trade_rets, trade_log, curve, dates, closes)


# ── 数据下载 + 六策略调用 ─────────────────────────────────────────────────

def _load_crypto_candles(ticker, timeframe):
    root = Path(__file__).resolve().parents[2]
    sys.path.insert(0, str(root / "backend"))
    from app.providers.local_crypto import LocalCryptoProvider

    candles = asyncio.run(LocalCryptoProvider().get_full_history(ticker, timeframe))
    if not candles:
        return None
    return {
        "dates": [c["time"] for c in candles],
        "opens": [c["open"] for c in candles],
        "highs": [c["high"] for c in candles],
        "lows": [c["low"] for c in candles],
        "closes": [c["close"] for c in candles],
    }


def _load_yfinance_daily(ticker):
    df = yf.download(ticker, period="max", auto_adjust=True, progress=False)
    if df.empty:
        return None
    return {
        "dates": df.index.tolist(),
        "opens": df["Open"].squeeze().tolist(),
        "highs": df["High"].squeeze().tolist(),
        "lows": df["Low"].squeeze().tolist(),
        "closes": df["Close"].squeeze().tolist(),
    }


def backtest_symbol_3x(ticker, timeframe="1day"):
    print(f"  下载数据: {ticker} {timeframe} ...", end=" ", flush=True)
    if timeframe == "1day":
        ohlc_src = _load_yfinance_daily(ticker)
    elif timeframe == "4h" and ticker in CRYPTO_SYMBOLS:
        ohlc_src = _load_crypto_candles(ticker, timeframe)
    else:
        print("该周期无数据源，跳过")
        return None
    if not ohlc_src or len(ohlc_src["closes"]) < max(MA_PERIODS) + 10:
        print("无数据，跳过")
        return None
    closes = ohlc_src["closes"]
    highs  = ohlc_src["highs"]
    lows   = ohlc_src["lows"]
    opens  = ohlc_src["opens"]
    dates  = ohlc_src["dates"]
    print(f"{len(closes)} 条K线  ({str(dates[0])[:10]} ~ {str(dates[-1])[:10]})")

    trend_sma = rolling_sma(closes, TREND_MA)
    rsi_vals  = rolling_rsi(closes)
    sma_by_n  = {n: rolling_sma(closes, n) for n in MA_PERIODS}

    out = {k: [] for k in [
        "breakout_1x", "pullback_1x",
        "breakout_2x", "pullback_2x",
        "breakout_3x", "pullback_3x",
    ]}

    for n in MA_PERIODS:
        # 1x 基准（使用原策略，保本 4%）
        r = run_breakout(closes, lows, opens, dates, n)
        if r: out["breakout_1x"].append(r)
        r = run_pullback(closes, lows, opens, dates, n, trend_sma, rsi_vals)
        if r: out["pullback_1x"].append(r)

        # 2x / 3x
        for lev in LEVERAGE_LEVELS:
            key = f"breakout_{lev}x"
            r = run_leveraged_breakout(closes, lows, opens, dates, n, lev)
            if r: out[key].append(r)
            key = f"pullback_{lev}x"
            r = run_leveraged_pullback(closes, lows, opens, dates, n, trend_sma, rsi_vals, lev)
            if r: out[key].append(r)

    # 各组最优 MA
    for key in list(out.keys()):
        vals = out[key]
        out[f"{key}_best"] = max(vals, key=lambda x: x["total_return"]) if vals else None

    out["ohlc"] = {
        "dates": dates, "opens": opens, "highs": highs,
        "lows": lows, "closes": closes,
        "sma_by_n": sma_by_n, "trend_sma": trend_sma,
    }
    return out


# ── Excel 样式 ───────────────────────────────────────────────────────────────

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
SUBHDR_FILL = PatternFill("solid", fgColor="2E75B6")
BEST_FILL   = PatternFill("solid", fgColor="FFD966")
ALT_FILL    = PatternFill("solid", fgColor="F5F5F5")
WHITE_FILL  = PatternFill("solid", fgColor="FFFFFF")
WIN_FILL    = PatternFill("solid", fgColor="E2EFDA")
LOSS_FILL   = PatternFill("solid", fgColor="FCE4D6")
ADD_FILL    = PatternFill("solid", fgColor="FFF3CD")
thin        = Side(style="thin", color="CCCCCC")
BORDER      = Border(left=thin, right=thin, top=thin, bottom=thin)

LEV_COLOR = {1: "1F3864", 2: "0070C0", 3: "C00000"}


def _cell(ws, row, col, value=None, fmt=None, bold=False, size=9,
          color="000000", fill=None, align="center", wrap=False):
    cell            = ws.cell(row, col)
    cell.value      = value
    cell.font       = Font(name="Arial", bold=bold, size=size, color=color)
    cell.alignment  = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    cell.border     = BORDER
    if fill: cell.fill = fill
    if fmt:  cell.number_format = fmt
    return cell


def _hdr(ws, row, col, text, fill=HEADER_FILL, color="FFFFFF", size=9):
    return _cell(ws, row, col, text, bold=True, color=color, fill=fill, size=size)


def color_ret(val):
    if val > 0: return "375623"
    if val < 0: return "C00000"
    return "000000"


def fmt_date(d):
    return d.strftime("%Y-%m-%d") if hasattr(d, "strftime") else str(d)[:10]


# ── Sheet 1：总对比（六组最优 MA 汇总） ─────────────────────────────────────

COMPARE_HEADERS = [
    "品种代码", "品种名称",
    "A-1x 最优MA", "A-1x 收益%", "A-1x 回撤%", "A-1x 胜率%",
    "A-2x 最优MA", "A-2x 收益%", "A-2x 回撤%", "A-2x 胜率%",
    "A-3x 最优MA", "A-3x 收益%", "A-3x 回撤%", "A-3x 胜率%",
    "B-1x 最优MA", "B-1x 收益%", "B-1x 回撤%", "B-1x 胜率%",
    "B-2x 最优MA", "B-2x 收益%", "B-2x 回撤%", "B-2x 胜率%",
    "B-3x 最优MA", "B-3x 收益%", "B-3x 回撤%", "B-3x 胜率%",
]


def write_compare_sheet(wb, all_results):
    ws = wb.create_sheet("总对比")
    ws.merge_cells(f"A1:{get_column_letter(len(COMPARE_HEADERS))}1")
    c = ws["A1"]
    c.value = (f"杠杆浮盈加仓策略对比  |  手续费 {FEE_RATE*100:.2f}%  "
               f"滑点 {SLIPPAGE*100:.2f}%  |  截至 {END_DATE}  "
               f"|  1x 保本4% / 2x-3x 保本15%")
    c.font  = Font(name="Arial", bold=True, size=11, color="FFFFFF")
    c.fill  = HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    for col, h in enumerate(COMPARE_HEADERS, 1):
        _hdr(ws, 2, col, h, fill=SUBHDR_FILL)
    ws.row_dimensions[2].height = 22

    items = list(all_results.items())
    items.sort(key=lambda kv: (
        ((kv[1] or {}).get("breakout_1x_best") or {})
    ).get("total_return", -1e9), reverse=True)

    for rank, (sym, data) in enumerate(items, 1):
        row  = rank + 2
        fill = ALT_FILL if rank % 2 == 0 else WHITE_FILL
        name = SYMBOLS.get(sym, sym)
        _cell(ws, row, 1, sym, fill=fill)
        _cell(ws, row, 2, name, fill=fill)
        col = 3
        for strat in ("breakout", "pullback"):
            for lev in (1, 2, 3):
                b = (data or {}).get(f"{strat}_{lev}x_best")
                if b:
                    ret_color = color_ret(b["total_return"])
                    _cell(ws, row, col,     b["n"],              fmt="0",         bold=True, fill=fill)
                    _cell(ws, row, col + 1, b["total_return"],   fmt='0.00"%"',   color=ret_color, bold=True, fill=fill)
                    _cell(ws, row, col + 2, b["max_drawdown"],   fmt='0.00"%"',   color="C00000" if b["max_drawdown"] > 100 else "000000", fill=fill)
                    _cell(ws, row, col + 3, b["win_rate"],       fmt='0.0"%"',    fill=fill)
                else:
                    for c_ in range(col, col + 4):
                        _cell(ws, row, c_, "-", fill=fill)
                col += 4

    col_widths = [12, 14] + [9, 12, 10, 9] * 6
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"


# ── Sheet 2-3：各策略全 MA × 全 leverage 汇总 ─────────────────────────────

SUMMARY_COLS = [
    "品种", "MA", "杠杆",
    "总收益%", "交易笔", "最大回撤%", "胜率%", "最大单笔%", "平均盈利%", "平均RR", "是否最优"
]


def write_strategy_sheet(wb, all_results, strat_key, strat_label):
    ws = wb.create_sheet(strat_label)
    ws.merge_cells(f"A1:{get_column_letter(len(SUMMARY_COLS))}1")
    c = ws["A1"]
    c.value = (f"{strat_label}  |  MA {MA_PERIODS}  |  手续费 {FEE_RATE*100:.2f}%  "
               f"滑点 {SLIPPAGE*100:.2f}%  |  截至 {END_DATE}")
    c.font  = Font(name="Arial", bold=True, size=11, color="FFFFFF")
    c.fill  = HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    for col, h in enumerate(SUMMARY_COLS, 1):
        _hdr(ws, 2, col, h, fill=SUBHDR_FILL)
    ws.row_dimensions[2].height = 22

    row = 3
    for sym, data in all_results.items():
        name = SYMBOLS.get(sym, sym)
        sym_start = row
        for lev in (1, 2, 3):
            key     = f"{strat_key}_{lev}x"
            results = (data or {}).get(key) or []
            best    = (data or {}).get(f"{key}_best")
            lev_fill = PatternFill("solid", fgColor=("E8F0FE" if lev == 2 else ("FCE4D6" if lev == 3 else "F5F5F5")))

            for r in results:
                is_best = best is not None and r["n"] == best["n"]
                fill    = BEST_FILL if is_best else lev_fill
                _cell(ws, row, 1, f"{sym}\n({name})", align="left", fill=fill, wrap=True)
                _cell(ws, row, 2, r["n"],           fmt="0",       fill=fill)
                _cell(ws, row, 3, f"{lev}x",                       fill=fill, bold=(lev > 1))
                _cell(ws, row, 4, r["total_return"],fmt='0.00"%"', color=color_ret(r["total_return"]), bold=is_best, fill=fill)
                _cell(ws, row, 5, r["n_trades"],    fmt="0",       fill=fill)
                _cell(ws, row, 6, r["max_drawdown"],fmt='0.00"%"', color="C00000" if r["max_drawdown"] > 100 else "000000", fill=fill)
                _cell(ws, row, 7, r["win_rate"],    fmt='0.0"%"',  fill=fill)
                _cell(ws, row, 8, r["max_profit"],  fmt='0.00"%"', color=color_ret(r["max_profit"]), fill=fill)
                _cell(ws, row, 9, r["avg_profit"],  fmt='0.00"%"', color=color_ret(r["avg_profit"]), fill=fill)
                avg_rr = r.get("avg_rr")
                _cell(ws, row, 10, round(avg_rr, 2) if avg_rr else "-",
                      fmt="0.00" if avg_rr else None, fill=fill)
                _cell(ws, row, 11, "★ 最优" if is_best else "",
                      bold=is_best, color="375623" if is_best else "000000", fill=fill)
                row += 1

        if row > sym_start + 1:
            ws.merge_cells(f"A{sym_start}:A{row - 1}")
            ws[f"A{sym_start}"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    col_widths = [16, 7, 7, 12, 9, 12, 9, 12, 12, 9, 10]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"


# ── Sheet 4-7：逐笔明细（含加仓事件子行，仅最优 MA × 2x/3x） ─────────────

TRADE_COLS = ["入场日期", "出场日期", "原因", "事件", "价格", "加仓量", "单笔收益%", "最大盈利%", "RR", "盈亏"]


def write_trade_detail_sheet(wb, all_results, strat_key, lev, sheet_label):
    ws = wb.create_sheet(sheet_label)
    ws.merge_cells(f"A1:{get_column_letter(len(TRADE_COLS))}1")
    c = ws["A1"]
    c.value = (f"{sheet_label}  |  杠杆 {lev}x  保本触发 +15%  "
               f"|  手续费 {FEE_RATE*100:.2f}%  滑点 {SLIPPAGE*100:.2f}%  |  截至 {END_DATE}")
    c.font  = Font(name="Arial", bold=True, size=11, color="FFFFFF")
    c.fill  = HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    cur = 2
    key = f"{strat_key}_{lev}x"
    best_key = f"{key}_best"

    for sym, data in all_results.items():
        if not data:
            continue
        best = data.get(best_key)
        if not best:
            continue
        results = data.get(key) or []
        best_n  = best["n"]
        # 只展示最优 MA 的逐笔
        r = next((x for x in results if x["n"] == best_n), None)
        if not r or not r.get("trade_log"):
            continue

        dates  = r["dates"]
        closes = r["closes"]
        name   = SYMBOLS.get(sym, sym)

        # 品种标题行
        ws.merge_cells(f"A{cur}:{get_column_letter(len(TRADE_COLS))}{cur}")
        tc = ws.cell(cur, 1)
        tc.value = (f"{sym} ({name})  ·  最优 MA={best_n}  ·  "
                    f"总收益 {r['total_return']:.2f}%  ·  {r['n_trades']} 笔  "
                    f"·  胜率 {r['win_rate']:.1f}%  ·  最大回撤 {r['max_drawdown']:.1f}%")
        tc.font  = Font(name="Arial", bold=True, size=10, color="FFFFFF")
        tc.fill  = PatternFill("solid", fgColor=("0070C0" if lev == 2 else "C00000"))
        tc.alignment = Alignment(horizontal="left", vertical="center")
        tc.border = BORDER
        ws.row_dimensions[cur].height = 20
        cur += 1

        for col, h in enumerate(TRADE_COLS, 1):
            _hdr(ws, cur, col, h, fill=SUBHDR_FILL, size=8)
        ws.row_dimensions[cur].height = 18
        cur += 1

        for tlog in r["trade_log"]:
            ei, xi, ret, reason, sl, rr, max_ret, _lev, events = tlog
            trade_fill = WIN_FILL if ret > 0 else LOSS_FILL
            entry_d    = fmt_date(dates[ei])
            exit_d     = fmt_date(dates[xi])

            # 主行（入场事件）
            _cell(ws, cur, 1, entry_d, fill=trade_fill, size=8)
            _cell(ws, cur, 2, exit_d,  fill=trade_fill, size=8)
            _cell(ws, cur, 3, reason,  fill=trade_fill, size=8,
                  color="C00000" if reason == "止损" else "D06000" if reason == "保本" else "000000")
            _cell(ws, cur, 4, "入场",  fill=trade_fill, size=8)
            ep = events[0]["price"] if events else closes[ei] * (1 + COST)
            _cell(ws, cur, 5, round(ep, 4), fmt="#,##0.####", fill=trade_fill, size=8)
            _cell(ws, cur, 6, float(lev),   fmt="0.00", fill=trade_fill, size=8)
            _cell(ws, cur, 7, ret * 100,    fmt='0.00"%"', color=color_ret(ret), bold=True, fill=trade_fill, size=8)
            _cell(ws, cur, 8, max_ret * 100, fmt='0.00"%"', color="375623", fill=trade_fill, size=8)
            _cell(ws, cur, 9, round(rr, 2) if rr else "-", fmt="0.00" if rr else None, fill=trade_fill, size=8)
            _cell(ws, cur, 10, "盈" if ret > 0 else "亏",
                  color="375623" if ret > 0 else "C00000", bold=True, fill=trade_fill, size=8)
            cur += 1

            # 加仓子行
            adds = [e for e in events if e["type"] == "add"]
            for ev in adds:
                add_d = fmt_date(dates[ev["bar_idx"]])
                _cell(ws, cur, 1, "",    fill=ADD_FILL, size=8)
                _cell(ws, cur, 2, add_d, fill=ADD_FILL, size=8, color="7B6000")
                _cell(ws, cur, 3, "",    fill=ADD_FILL, size=8)
                _cell(ws, cur, 4, "加仓", fill=ADD_FILL, size=8, color="7B6000", bold=True)
                _cell(ws, cur, 5, round(ev["price"], 4), fmt="#,##0.####", fill=ADD_FILL, size=8)
                _cell(ws, cur, 6, round(ev["size"], 4),  fmt="0.0000",    fill=ADD_FILL, size=8, color="7B6000")
                for c_ in range(7, 11):
                    _cell(ws, cur, c_, "", fill=ADD_FILL, size=8)
                cur += 1

        cur += 1  # 品种间留空行

    col_widths = [12, 12, 9, 8, 14, 10, 12, 12, 9, 8]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


# ── 主流程 ───────────────────────────────────────────────────────────────────

def main(only_symbol=None, only_leverage=None, skip_excel=False, timeframe="1day"):
    if timeframe not in ("1day", "4h"):
        raise SystemExit("未知周期：仅支持 1day / 4h")
    targets = [only_symbol] if only_symbol else (list(CRYPTO_SYMBOLS) if timeframe == "4h" else list(SYMBOLS.keys()))
    if only_symbol and only_symbol not in SYMBOLS:
        raise SystemExit(f"未知品种：{only_symbol}（合法范围：{', '.join(SYMBOLS)})")

    all_results = {}
    for sym in targets:
        print(f"\n[{sym}]")
        data = backtest_symbol_3x(sym, timeframe=timeframe)
        all_results[sym] = data
        if data:
            for strat in ("breakout", "pullback"):
                strat_label = "突破A" if strat == "breakout" else "回调B"
                for lev in (1, 2, 3):
                    b = data.get(f"{strat}_{lev}x_best")
                    if b:
                        rr_str = f"  RR={b['avg_rr']:.2f}" if b.get("avg_rr") else ""
                        adds_k = sum(
                            len([e for e in t[8] if e["type"] == "add"]) if len(t) > 8 else 0
                            for t in b.get("trade_log", [])
                        )
                        print(f"  [{strat_label} {lev}x] MA={b['n']}  "
                              f"收益={b['total_return']:.1f}%  "
                              f"胜率={b['win_rate']:.1f}%  "
                              f"回撤={b['max_drawdown']:.1f}%"
                              f"{rr_str}"
                              + (f"  共加仓{adds_k}次" if lev > 1 else ""))
                    else:
                        print(f"  [{strat_label} {lev}x] 无有效交易")

    # ── JSON 导出（供前端读取） ─────────────────────────────────────────────
    from backtest_json_export_3x import export_json_3x, write_manifest_3x
    for lev in (LEVERAGE_LEVELS if not only_leverage else [only_leverage]):
        dir_name = f"ma_strategies_{lev}x" if timeframe == "1day" else f"ma_strategies_{lev}x_{timeframe}"
        json_dir = str(_REPO.parent.parent / "backend" / "data" / dir_name)
        os.makedirs(json_dir, exist_ok=True)
        print(f"\n导出 {lev}x JSON → {json_dir}")
        for sym, data in all_results.items():
            if not data or "ohlc" not in data:
                continue
            safe    = sym.replace("/", "-")
            best_b  = (data.get(f"breakout_{lev}x_best") or {}).get("n")
            best_p  = (data.get(f"pullback_{lev}x_best") or {}).get("n")
            export_json_3x(
                sym, SYMBOLS.get(sym, sym), data["ohlc"],
                data.get(f"breakout_{lev}x") or [],
                data.get(f"pullback_{lev}x") or [],
                os.path.join(json_dir, f"{safe}.json"),
                breakout_best=best_b, pullback_best=best_p,
                leverage=lev,
            )
            print(f"  ✓ {sym}.json")
        if not only_symbol:
            write_manifest_3x(SYMBOLS, json_dir)
        print(f"✅ 已导出 {lev}x JSON")

    # ── Excel（全量才输出） ────────────────────────────────────────────────
    if skip_excel or only_symbol or timeframe != "1day":
        return

    out_path = str(_REPO / "data" / "ma_strategies_3x_backtest.xlsx")
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    print("\n生成 Excel ...")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    write_compare_sheet(wb, all_results)
    write_strategy_sheet(wb, all_results, "breakout", "A-突破策略")
    write_strategy_sheet(wb, all_results, "pullback",  "B-回调策略")
    for strat_key, strat_label in [("breakout", "A-突破"), ("pullback", "B-回调")]:
        for lev in LEVERAGE_LEVELS:
            write_trade_detail_sheet(wb, all_results, strat_key, lev,
                                     f"{strat_label}-{lev}x明细")

    wb.save(out_path)
    print(f"✅ 已保存至 {out_path}")


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("--symbol",   help="只回测该品种")
    parser.add_argument("--leverage", type=int, choices=[2, 3], help="只导出指定杠杆的 JSON")
    parser.add_argument("--timeframe", choices=["1day", "4h"], default="1day", help="回测周期")
    parser.add_argument("--no-excel", action="store_true")
    args = parser.parse_args()
    main(only_symbol=args.symbol, only_leverage=args.leverage, skip_excel=args.no_excel, timeframe=args.timeframe)
