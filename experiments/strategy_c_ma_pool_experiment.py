"""
【实验·不动主回测系统】策略 C 入场均线池对比

对比两套「同时上穿池内两根均线」（即 C-x2 形态）的均线参数：
  - 池 A：MA21 + MA55
  - 池 B：MA30 + MA60
入场：当根 K 线同时上穿池内两根均线 + 实体涨幅 > 3%。
出场 / 止损 / 保本：完全沿用策略 C（= 策略 A）逻辑，MA_exit = 被上穿两根中周期较小那根。

品种：BTC-USD / ETH-USD    周期：1day / 4h / 1h
产出：experiments/strategy_c_ma_pool_compare.xlsx（独立文件，不写 JSON、不刷 manifest）。

复用 scripts/backtests/backtest_ma_strategies_c.py 的核心函数，仅把固定的 4 均线池
换成可传参的 2 均线池（run_variant 本身与均线池无关）。
"""

import datetime
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

# ── 让本实验能 import 主回测脚本 ─────────────────────────────────────────────
_REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(_REPO / "scripts" / "backtests"))

# 导入 C 脚本会自动把 backend / scripts 目录挂到 sys.path 上
import backtest_ma_strategies_c as C  # noqa: E402
from backtest_ma_strategies_c import (  # noqa: E402
    load_candles, run_variant, summarize,
    COST, FEE_RATE, SLIPPAGE, BODY_RISE_MIN, MA_EXIT_PCT,
    _cell, _hdr, color_ret,
    HEADER_FILL, SUBHDR_FILL, BEST_FILL, SEC_FILL, COL_FILL,
    ALT_FILL, WHITE_FILL, WIN_FILL, LOSS_FILL,
)
from backtest_ma_strategies import rolling_sma  # noqa: E402

# ── 实验配置 ─────────────────────────────────────────────────────────────────
SYMBOLS = {"BTC-USD": "BTC", "ETH-USD": "ETH"}
TIMEFRAMES = ["1day", "4h", "1h"]
POOLS = {
    "A": [21, 55],
    "B": [30, 60],
}
END_DATE = datetime.date.today().isoformat()
XLSX_PATH = _REPO / "experiments" / "strategy_c_ma_pool_compare.xlsx"


def pool_label(key):
    return "MA" + "+MA".join(str(m) for m in POOLS[key])  # e.g. "MA21+MA55"


def compute_crossed_by_bar(closes, sma_by_n, pool):
    """每根 K 线上穿了池内哪几根均线（prev_c ≤ prev_ma 且 c > ma）。"""
    n = len(closes)
    out = [[] for _ in range(n)]
    for i in range(1, n):
        c, pc = closes[i], closes[i - 1]
        crossed = []
        for m in pool:
            ma, pma = sma_by_n[m][i], sma_by_n[m][i - 1]
            if ma is not None and pma is not None and pc <= pma and c > ma:
                crossed.append(m)
        out[i] = crossed
    return out


def backtest(symbol, tf, pool):
    candles = load_candles(symbol, tf)
    if not candles or len(candles) < max(pool) + 10:
        return None
    times = [c["time"] for c in candles]
    opens = [c["open"] for c in candles]
    lows = [c["low"] for c in candles]
    closes = [c["close"] for c in candles]

    sma_by_n = {m: rolling_sma(closes, m) for m in pool}
    crossed_by_bar = compute_crossed_by_bar(closes, sma_by_n, pool)
    # 同时上穿全部两根 → target_k = 2
    r = run_variant(closes, opens, lows, times, sma_by_n, crossed_by_bar, target_k=len(pool))
    if r:
        r["_times"] = times
        r["_closes"] = closes
    return r


# ── Excel ────────────────────────────────────────────────────────────────────

def _fmt_dt(unix, tf):
    d = datetime.datetime.utcfromtimestamp(int(unix))
    return d.strftime("%Y-%m-%d") if tf == "1day" else d.strftime("%Y-%m-%d %H:%M")


def write_compare_sheet(wb, results):
    ws = wb.create_sheet("对比汇总")
    ws.merge_cells("A1:M1")
    c = ws["A1"]
    c.value = (f"策略C 入场均线池对比（同时上穿池内两根均线 + 实体涨幅>{BODY_RISE_MIN*100:.0f}%，"
               f"均线出场≥{MA_EXIT_PCT*100:.0f}%）  |  手续费 {FEE_RATE*100:.2f}% 滑点 {SLIPPAGE*100:.2f}%  |  截至 {END_DATE}")
    c.font = Font(name="Arial", bold=True, size=12, color="FFFFFF")
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    la, lb = pool_label("A"), pool_label("B")
    headers = ["品种", "周期",
               f"A.{la} 收益%", "A.交易", "A.回撤%", "A.胜率%", "A.盈亏比R",
               f"B.{lb} 收益%", "B.交易", "B.回撤%", "B.胜率%", "B.盈亏比R",
               "更优"]
    for col, h in enumerate(headers, 1):
        _hdr(ws, 2, col, h, fill=SUBHDR_FILL)
    ws.row_dimensions[2].height = 22

    row = 3
    for symbol in SYMBOLS:
        name = SYMBOLS[symbol]
        start_row = row
        for tf in TIMEFRAMES:
            ra = results.get((symbol, tf, "A"))
            rb = results.get((symbol, tf, "B"))
            fill = ALT_FILL if row % 2 == 0 else WHITE_FILL
            _cell(ws, row, 1, f"{symbol}\n({name})", align="left", fill=fill)
            _cell(ws, row, 2, tf, fill=fill)

            def put(base_col, r):
                if r:
                    _cell(ws, row, base_col, r["total_return"], fmt='0.00"%"',
                          color=color_ret(r["total_return"]), bold=True, fill=fill)
                    _cell(ws, row, base_col + 1, r["n_trades"], fmt="0", fill=fill)
                    _cell(ws, row, base_col + 2, r["max_drawdown"], fmt='0.00"%"',
                          color="C00000" if r["max_drawdown"] > 50 else "000000", fill=fill)
                    _cell(ws, row, base_col + 3, r["win_rate"], fmt='0.0"%"', fill=fill)
                    rr = r["avg_rr"]
                    _cell(ws, row, base_col + 4, round(rr, 2) if rr is not None else "-",
                          fmt="0.00" if rr is not None else None, fill=fill)
                else:
                    for k in range(5):
                        _cell(ws, row, base_col + k, "-", fill=fill)

            put(3, ra)
            put(8, rb)

            ar = ra["total_return"] if ra else float("-inf")
            br = rb["total_return"] if rb else float("-inf")
            winner = "A" if ar > br else ("B" if br > ar else "=")
            _cell(ws, row, 13, winner, bold=True,
                  color="375623" if winner == "A" else ("1F3864" if winner == "B" else "000000"), fill=fill)
            row += 1
        if row - start_row > 1:
            ws.merge_cells(f"A{start_row}:A{row-1}")
            ws[f"A{start_row}"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    widths = [16, 8, 15, 9, 11, 10, 11, 15, 9, 11, 10, 11, 8]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"


TRADE_COLS = ["入场日期", "出场日期", "出场原因", "入场有效价", "出场有效价",
              "单笔收益率 %", "最大盈利率 %", "盈亏比 R", "出场参考线 MA", "上穿均线", "盈亏"]
N_COLS = len(TRADE_COLS)


def write_trade_sheet(wb, symbol, tf, results):
    name = SYMBOLS[symbol]
    ws = wb.create_sheet(f"{name}-{tf}"[:28])

    ws.merge_cells(f"A1:{get_column_letter(N_COLS)}1")
    c = ws["A1"]
    c.value = f"策略C 均线池对比 | {symbol} ({name}) · {tf}  —  手续费 {FEE_RATE*100:.2f}% 滑点 {SLIPPAGE*100:.2f}%"
    c.font = Font(name="Arial", bold=True, size=12, color="FFFFFF")
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    cur = 2
    # 哪个池更优 → 段标题加 ★
    rA = results.get((symbol, tf, "A"))
    rB = results.get((symbol, tf, "B"))
    best_key = None
    if rA or rB:
        best_key = "A" if (rA["total_return"] if rA else -1e18) >= (rB["total_return"] if rB else -1e18) else "B"

    for key in ("A", "B"):
        r = results.get((symbol, tf, key))
        is_best = (key == best_key)
        sec_fill = BEST_FILL if is_best else SEC_FILL
        sec_color = "8B4513" if is_best else "FFFFFF"
        ws.merge_cells(f"A{cur}:{get_column_letter(N_COLS)}{cur}")
        tc = ws.cell(cur, 1)
        prefix = f"{'★ 更优  ' if is_best else ''}池{key} = {pool_label(key)}  |  "
        if r:
            base = (f"总收益 {r['total_return']:.2f}%  |  交易 {r['n_trades']} 笔  |  "
                    f"胜率 {r['win_rate']:.1f}%  |  最大回撤 {r['max_drawdown']:.1f}%  |  "
                    f"最大单笔盈利 {r['max_profit']:.2f}%")
            if r.get("avg_rr"):
                base += f"  |  平均盈亏比 {r['avg_rr']:.2f}R"
        else:
            base = "无有效交易"
        tc.value = prefix + base
        tc.font = Font(name="Arial", bold=True, size=10, color=sec_color)
        tc.fill = sec_fill
        tc.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[cur].height = 20
        cur += 1

        for col, h in enumerate(TRADE_COLS, 1):
            _hdr(ws, cur, col, h, fill=COL_FILL, size=9)
        ws.row_dimensions[cur].height = 18
        cur += 1

        trades = r["trade_log"] if r else []
        if not trades:
            ws.merge_cells(f"A{cur}:{get_column_letter(N_COLS)}{cur}")
            _cell(ws, cur, 1, "该参数下无交易记录", color="888888", fill=WHITE_FILL, align="left")
            cur += 1
        else:
            times = r["_times"]
            closes = r["_closes"]
            for (ei, xi, ret, reason, sl, rr, max_ret, exit_ma, crossed) in trades:
                fill = WIN_FILL if ret > 0 else LOSS_FILL
                entry_px = closes[ei] * (1 + COST)
                exit_px = sl * (1 - COST) if reason in ("止损", "保本") else closes[xi] * (1 - COST)
                reason_color = ("C00000" if reason == "止损" else "D06000" if reason == "保本" else "000000")
                _cell(ws, cur, 1, _fmt_dt(times[ei], tf), fill=fill)
                _cell(ws, cur, 2, _fmt_dt(times[xi], tf), fill=fill)
                _cell(ws, cur, 3, reason, fill=fill, color=reason_color)
                _cell(ws, cur, 4, round(entry_px, 6), fmt="#,##0.######", fill=fill)
                _cell(ws, cur, 5, round(exit_px, 6), fmt="#,##0.######", fill=fill)
                _cell(ws, cur, 6, ret * 100, fmt='0.00"%"', color=color_ret(ret), bold=(ret > 0), fill=fill)
                _cell(ws, cur, 7, max_ret * 100, fmt='0.00"%"', color="375623", bold=True, fill=fill)
                if rr is not None:
                    _cell(ws, cur, 8, round(rr, 2), fmt="0.00", color="375623", bold=True, fill=fill)
                else:
                    _cell(ws, cur, 8, "-", color="888888", fill=fill)
                _cell(ws, cur, 9, f"MA{exit_ma}", fill=fill)
                _cell(ws, cur, 10, "/".join(f"MA{m}" for m in crossed), fill=fill)
                _cell(ws, cur, 11, "盈" if ret > 0 else "亏",
                      color="375623" if ret > 0 else "C00000", bold=True, fill=fill)
                cur += 1
        cur += 1

    widths = [18, 18, 10, 16, 16, 14, 14, 10, 14, 16, 8]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


# ── 主流程 ───────────────────────────────────────────────────────────────────

def main():
    results = {}
    for symbol in SYMBOLS:
        print(f"\n[{symbol}]")
        for tf in TIMEFRAMES:
            print(f"  回测 {tf} ...", flush=True)
            for key, pool in POOLS.items():
                r = backtest(symbol, tf, pool)
                results[(symbol, tf, key)] = r
                if r:
                    rr = f"  盈亏比={r['avg_rr']:.2f}R" if r["avg_rr"] else ""
                    print(f"    池{key} {pool_label(key)}: 收益={r['total_return']:8.1f}%  "
                          f"交易={r['n_trades']:4d}  胜率={r['win_rate']:5.1f}%  "
                          f"回撤={r['max_drawdown']:5.1f}%{rr}")
                else:
                    print(f"    池{key} {pool_label(key)}: 数据不足/无交易")

    print("\n生成 Excel ...")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    write_compare_sheet(wb, results)
    for symbol in SYMBOLS:
        for tf in TIMEFRAMES:
            write_trade_sheet(wb, symbol, tf, results)
    wb.save(str(XLSX_PATH))
    print(f"✅ 已保存实验 Excel 至 {XLSX_PATH}")


if __name__ == "__main__":
    main()
